from openpyxl import load_workbook

from imdbtop.__main__ import main
from imdbtop.build import build_top
from imdbtop.export import export
from imdbtop.fetch import find_dataset_links

INDEX_HTML = """
<html><body><ul>
<li><a href="https://datasets.imdbws.com/name.basics.tsv.gz">name.basics.tsv.gz</a></li>
<li><a href="title.ratings.tsv.gz">title.ratings.tsv.gz</a></li>
<li><a href="/title.basics.tsv.gz ">title.basics.tsv.gz</a></li>
<li><a href="https://developer.imdb.com/">docs</a></li>
</ul></body></html>
"""


def test_find_dataset_links_resolves_relative_urls():
    links = find_dataset_links(INDEX_HTML)
    assert links == {
        "name.basics": "https://datasets.imdbws.com/name.basics.tsv.gz",
        "title.ratings": "https://datasets.imdbws.com/title.ratings.tsv.gz",
        "title.basics": "https://datasets.imdbws.com/title.basics.tsv.gz",
    }


def test_workbook_layout(dataset, tmp_path):
    movies = build_top(dataset, log=lambda *_: None)
    out = export(movies, tmp_path / "out" / "top.xlsx")
    wb = load_workbook(out)

    assert wb.sheetnames == ["Top 250", "Summary", "By Decade", "By Genre", "Directors", "Runtime", "About"]

    ws = wb["Top 250"]
    assert ws["A1"].value == "Rank" and ws["B2"].value == "Harbour Lights"
    assert ws.max_row == len(movies) + 1
    assert ws["L2"].hyperlink.target == "https://www.imdb.com/title/tt0000001/"
    assert ws.freeze_panes == "C2"
    # Excel rejects table names that look like cell addresses (e.g. TOP250).
    assert list(ws.tables) == ["Movies"]

    summary = wb["Summary"]
    assert summary["B4"].value == "=COUNT('Top 250'!I2:I6)"
    formulas = [c.value for (c,) in summary.iter_rows(min_row=4, min_col=2, max_col=2) if isinstance(c.value, str)]
    # ">=8.5" inside a string breaks in locales that use a decimal comma.
    assert not any('"' in f and "." in f.split('"')[1] for f in formulas if f.startswith("=COUNTIF"))

    assert wb["By Decade"]["A2"].value == 1950
    assert "Made by Talha Aker" in [r[0] for r in wb["About"].iter_rows(values_only=True)]


def test_cli_offline(dataset, tmp_path, capsys):
    data_dir = dataset["title.ratings"].parent
    out = tmp_path / "cli.xlsx"
    charts = tmp_path / "charts"
    assert main(["--data-dir", str(data_dir), "--out", str(out), "--offline", "--charts", str(charts)]) == 0
    assert out.exists()
    assert (charts / "decades.svg").read_text(encoding="utf-8").startswith("<svg")
    assert "1. Harbour Lights" in capsys.readouterr().out


def test_cli_offline_reports_missing_files(tmp_path):
    assert main(["--data-dir", str(tmp_path), "--offline"]) == 1


def test_site_shows_top_10_only(dataset, tmp_path):
    from imdbtop import site

    movies = build_top(dataset, log=lambda *_: None)
    html = site.write(movies * 3, tmp_path / "site").read_text(encoding="utf-8")
    assert html.count("<tr><td class=\"num\">") == 10
    assert "&quot;Quote&quot; Story" in html  # titles are HTML escaped
    assert "Made by Talha Aker" in html
