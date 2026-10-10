"""Write the Top 250 and the analysis sheets to an Excel workbook with OpenPyXL."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import analyze
from .build import MIN_VOTES, Movie

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)

# (header, width, number format, value getter)
TOP_COLUMNS = [
    ("Rank", 7, "0", lambda m: m.rank),
    ("Title", 42, None, lambda m: m.title),
    ("Original Title", 36, None, lambda m: m.original_title),
    ("Year", 7, "0", lambda m: m.year),
    ("Decade", 8, "0", lambda m: m.decade),
    ("Runtime (min)", 14, "0", lambda m: m.runtime),
    ("Genres", 28, None, lambda m: ", ".join(m.genres)),
    ("Director(s)", 34, None, lambda m: ", ".join(m.directors)),
    ("IMDb Rating", 12, "0.0", lambda m: m.rating),
    ("Votes", 12, "#,##0", lambda m: m.votes),
    ("Weighted Score", 15, "0.000", lambda m: round(m.score, 3)),
    ("IMDb Link", 38, None, lambda m: m.url),
]


def _col(header: str) -> str:
    return get_column_letter([c[0] for c in TOP_COLUMNS].index(header) + 1)


def _style_header(ws, row: int, ncols: int, start_col: int = 1) -> None:
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_top_sheet(ws, movies: list[Movie]) -> None:
    ws.title = "Top 250"
    ws.append([c[0] for c in TOP_COLUMNS])
    for m in movies:
        ws.append([get(m) for *_, get in TOP_COLUMNS])

    last = len(movies) + 1
    for i, (_, width, fmt, _) in enumerate(TOP_COLUMNS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=i, max_col=i):
                cell.number_format = fmt

    link_col = len(TOP_COLUMNS)
    for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=link_col, max_col=link_col):
        cell.hyperlink = cell.value
        cell.style = "Hyperlink"

    if movies:
        # A table gives sort and filter buttons on every column.
        ref = f"A1:{get_column_letter(len(TOP_COLUMNS))}{last}"
        table = Table(displayName="Movies", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        rating = _col("IMDb Rating")
        ws.conditional_formatting.add(
            f"{rating}2:{rating}{last}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
        )
    ws.freeze_panes = "C2"


def write_summary_sheet(ws, movies: list[Movie]) -> None:
    """Live Excel formulas over the Top 250 sheet, plus a few facts computed in Python."""
    last = len(movies) + 1
    rating = f"'Top 250'!{_col('IMDb Rating')}2:{_col('IMDb Rating')}{last}"
    votes = f"'Top 250'!{_col('Votes')}2:{_col('Votes')}{last}"
    runtime = f"'Top 250'!{_col('Runtime (min)')}2:{_col('Runtime (min)')}{last}"
    year = f"'Top 250'!{_col('Year')}2:{_col('Year')}{last}"

    ws["A1"] = "IMDb Top 250 - Summary"
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(["Metric", "Value"])
    _style_header(ws, 3, 2)

    facts = analyze.overview(movies)
    rows = [
        ("Movies", f"=COUNT({rating})", "0"),
        ("Average rating", f"=AVERAGE({rating})", "0.00"),
        ("Highest rating", f"=MAX({rating})", "0.0"),
        ("Lowest rating", f"=MIN({rating})", "0.0"),
        ("Total votes", f"=SUM({votes})", "#,##0"),
        ("Median votes", f"=MEDIAN({votes})", "#,##0"),
        ("Average runtime (min)", f"=AVERAGE({runtime})", "0"),
        ("Longest runtime (min)", f"=MAX({runtime})", "0"),
        ("Movies from 2000 or later", f'=COUNTIF({year},">=2000")', "0"),
        ("Movies before 1970", f'=COUNTIF({year},"<1970")', "0"),
        ("Rated 8.5 or higher", f'=COUNTIF({rating},">="&8.5)', "0"),
        ("Year span", facts["year_span"], None),
        ("Most voted", facts["most_voted"], None),
        ("Oldest", facts["oldest"], None),
        ("Newest", facts["newest"], None),
        ("Titles with a different original title", facts["translated_titles"], "0"),
    ]
    for label, value, fmt in rows:
        ws.append([label, value])
        if fmt:
            ws.cell(row=ws.max_row, column=2).number_format = fmt
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 48
    for (cell,) in ws.iter_rows(min_row=4, min_col=2, max_col=2):
        cell.alignment = Alignment(horizontal="left")


def write_group_sheet(ws, title: str, key_header: str, rows: list[dict]) -> int:
    """Write a grouped table and return the last data row."""
    ws.title = title
    headers = [key_header, "Movies", "Avg Rating", "Avg Votes", "Highest Ranked"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for r in rows:
        ws.append([r["key"], r["count"], r["avg_rating"], r["avg_votes"], r["best"]])
    for (cell,) in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        cell.number_format = "0.00"
    for (cell,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        cell.number_format = "#,##0"
    for letter, width in zip("ABCDE", (22, 9, 11, 12, 44)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:E{len(rows) + 1}"
    return len(rows) + 1


def _bar_chart(ws, title: str, last: int, anchor: str, x_title: str, horizontal: bool = False) -> None:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.y_axis.title = "Movies"
    chart.x_axis.title = x_title
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.width, chart.height = 18, 9 if not horizontal else 12
    ws.add_chart(chart, anchor)


def write_about_sheet(ws, movies: list[Movie]) -> None:
    ws.title = "About"
    lines = [
        ("IMDb Top 250 - data notes", TITLE_FONT),
        (f"Built on {date.today().isoformat()} from the IMDb non-commercial datasets (https://datasets.imdbws.com/).", None),
        (f"Feature films only, adult titles excluded, at least {MIN_VOTES:,} votes.", None),
        ("Ranked by weighted rating: WR = v/(v+m)*R + m/(v+m)*C, with m = 25,000 and C = mean rating of all eligible films.", None),
        ("IMDb's live chart uses an unpublished formula, so a few positions differ.", None),
        ("Data courtesy of IMDb (https://www.imdb.com). Personal and non-commercial use only.", None),
        ("", None),
        ("Made by Talha Aker", Font(italic=True)),
    ]
    for text, font in lines:
        ws.append([text])
        if font:
            ws.cell(row=ws.max_row, column=1).font = font
    ws.column_dimensions["A"].width = 110


def export(movies: list[Movie], out: Path) -> Path:
    wb = Workbook()
    write_top_sheet(wb.active, movies)
    write_summary_sheet(wb.create_sheet("Summary"), movies)

    ws = wb.create_sheet()
    last = write_group_sheet(ws, "By Decade", "Decade", analyze.by_decade(movies))
    _bar_chart(ws, "Top 250 movies per decade", last, "G2", "Decade")
    line = LineChart()
    line.title = "Average rating per decade"
    line.y_axis.title = "Rating"
    line.legend = None
    line.add_data(Reference(ws, min_col=3, min_row=1, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    line.width, line.height = 18, 9
    ws.add_chart(line, "G21")

    ws = wb.create_sheet()
    last = write_group_sheet(ws, "By Genre", "Genre", analyze.by_genre(movies))
    _bar_chart(ws, "Movies per genre (a movie can have up to 3)", last, "G2", "Genre", horizontal=True)

    ws = wb.create_sheet()
    last = write_group_sheet(ws, "Directors", "Director", analyze.by_director(movies))
    ws.column_dimensions["A"].width = 28

    ws = wb.create_sheet()
    last = write_group_sheet(ws, "Runtime", "Runtime", analyze.runtime_buckets(movies))
    _bar_chart(ws, "Movies by runtime", last, "G2", "Runtime")

    write_about_sheet(wb.create_sheet(), movies)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
